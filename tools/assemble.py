#!/usr/bin/env python3
"""Сборка DOCX из распознанных страниц.

python tools/assemble.py work/<имя> [--out output]

Берёт страницы из work/<имя>/plan.json -> work/<имя>/html/pNNN.html,
склеивает таблицы/строки/абзацы, разорванные между страницами,
собирает FLAG-комментарии и пишет:
  output/<имя>.docx        — документ
  output/<имя>_flags.md    — неразборчивые места и несходящиеся цифры
Конвертация через LibreOffice (если есть), иначе встроенный python-docx.
Модель не нужна.
"""
import argparse, json, re, shutil, subprocess, sys, tempfile
from pathlib import Path

try:
    from bs4 import BeautifulSoup, Comment, NavigableString, Tag
except ImportError:
    sys.exit("Нужен beautifulsoup4: pip install beautifulsoup4 python-docx")

TEMPLATE = """<!DOCTYPE html><html><head><meta charset="utf-8"><title>{title}</title>
<style>
body {{ font-family: 'Times New Roman', serif; font-size: 12pt; }}
p.caption {{ text-align: right; }}
span.unclear {{ background: yellow; }}
h1, h2, h3 {{ font-family: 'Times New Roman', serif; }}
h1 {{ font-size: 14pt; text-align: center; }}
h2 {{ font-size: 13pt; text-align: center; }}
h3 {{ font-size: 12pt; }}
</style></head><body>
{body}
</body></html>"""


def load_pages(work):
    plan = json.loads((work / "plan.json").read_text(encoding="utf-8"))
    missing = [p for p in plan["selected"] if not (work / "html" / f"p{p:03d}.html").exists()]
    return plan, missing


def last_top(body, name):
    kids = [k for k in body.children if isinstance(k, Tag)]
    if kids and kids[-1].name in (name if isinstance(name, tuple) else (name,)):
        return kids[-1]
    return None


def rows_of(table):
    return table.find_all("tr")


def merge(pages_html):
    master = BeautifulSoup("<div id='root'></div>", "html.parser")
    root = master.find(id="root")
    flags = []
    for page_id, html in pages_html:
        frag = BeautifulSoup(html, "html.parser")
        for c in frag.find_all(string=lambda s: isinstance(s, Comment)):
            text = c.strip()
            if text.startswith("FLAG"):
                msg = re.sub(r"^p\d{3}\s*[:\-]?\s*", "", text[4:].lstrip(": ").strip())
                flags.append(f"- {page_id}: {msg}")
            c.extract()
        for el in list(frag.children):
            if isinstance(el, NavigableString):
                if el.strip():
                    p = master.new_tag("p"); p.string = el.strip(); root.append(p)
                continue
            if el.name == "table" and el.get("data-continued") == "true":
                prev = root.find_all("table")[-1] if root.find_all("table") else None
                if prev is None:
                    flags.append(f"- {page_id}: продолжение таблицы без начала — вставлено как отдельная таблица")
                    del el["data-continued"]; root.append(el); continue
                new_rows = rows_of(el)
                target = prev.find("tbody") or prev
                if new_rows and new_rows[0].get("data-continues-row") == "true":
                    first = new_rows.pop(0)
                    last = rows_of(prev)[-1]
                    old_cells = last.find_all(["td", "th"], recursive=False)
                    new_cells = first.find_all(["td", "th"], recursive=False)
                    if len(old_cells) == len(new_cells):
                        for oc, nc in zip(old_cells, new_cells):
                            if nc.get_text(strip=True):
                                if oc.get_text(strip=True):
                                    oc.append(master.new_tag("br"))
                                for x in list(nc.contents):
                                    oc.append(x)
                    else:
                        del first["data-continues-row"]
                        target.append(first)
                        flags.append(f"- {page_id}: строку, разорванную между страницами, не удалось склеить автоматически — проверь вручную")
                for r in new_rows:
                    target.append(r)
                continue
            if el.name == "p" and el.get("data-continued") == "true":
                del el["data-continued"]
                prev = last_top(root, "p")
                if prev is None:
                    lst = last_top(root, ("ul", "ol"))
                    prev = lst.find_all("li")[-1] if lst and lst.find_all("li") else None
                if prev is not None:
                    prev.append(" ")
                    for x in list(el.contents):
                        prev.append(x)
                    continue
            root.append(el)
    for t in root.find_all("table"):
        t.attrs.pop("data-continued", None)
        t["border"] = "1"; t["cellspacing"] = "0"; t["cellpadding"] = "4"; t["width"] = "100%"
    for tr in root.find_all("tr"):
        tr.attrs.pop("data-continues-row", None)
    for sp in root.find_all("span", class_="unclear"):
        sp["style"] = "background:#ffff00"   # LibreOffice не берёт фон из CSS-класса
    return root.decode_contents(), flags


# ---------- запасной путь: python-docx ----------
def to_docx_fallback(body_html, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Pt

    doc = Document()
    from docx.shared import RGBColor
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    for lvl, size in ((1, 14), (2, 13), (3, 12)):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name = "Times New Roman"; hs.font.size = Pt(size); hs.font.color.rgb = RGBColor(0, 0, 0)
    from docx.oxml.ns import qn
    for sname in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        rf = doc.styles[sname].element.get_or_add_rPr().get_or_add_rFonts()
        for k in list(rf.attrib):
            if k.lower().endswith("theme"): del rf.attrib[k]
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), "Times New Roman")
    soup = BeautifulSoup(body_html, "html.parser")

    def runs(par, node, fmt):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                t = re.sub(r"\s+", " ", str(ch))
                if t:
                    r = par.add_run(t)
                    r.bold, r.italic, r.underline = fmt.get("b"), fmt.get("i"), fmt.get("u")
                    if fmt.get("hl"):
                        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif ch.name == "br":
                par.add_run().add_break()
            elif ch.name in ("b", "strong", "i", "em", "u", "span", "sup", "sub", "a"):
                f = dict(fmt)
                if ch.name in ("b", "strong"): f["b"] = True
                if ch.name in ("i", "em"): f["i"] = True
                if ch.name == "u": f["u"] = True
                if "unclear" in (ch.get("class") or []): f["hl"] = True
                runs(par, ch, f)
            else:
                runs(par, ch, fmt)

    def blocks(container, node, first_par=None):
        state = {"first": first_par}

        def newp(style=None):
            if state["first"] is not None:
                p = state["first"]; state["first"] = None
                if style: p.style = style
                return p
            return container.add_paragraph(style=style)

        inline_buf = []

        def flush():
            if inline_buf:
                p = newp()
                holder = soup.new_tag("span")
                for x in inline_buf: holder.append(x.__copy__() if isinstance(x, Tag) else NavigableString(str(x)))
                runs(p, holder, {})
                inline_buf.clear()

        for ch in list(node.children):
            if isinstance(ch, NavigableString) or (isinstance(ch, Tag) and ch.name in ("b","strong","i","em","u","span","br","sup","sub","a")):
                if isinstance(ch, NavigableString) and not ch.strip() and not inline_buf:
                    continue
                inline_buf.append(ch); continue
            flush()
            if ch.name in ("h1", "h2", "h3"):
                lvl = int(ch.name[1])
                if isinstance(container, type(doc)):
                    p = container.add_heading(level=lvl)
                else:
                    p = newp()
                runs(p, ch, {"b": True})
                if lvl <= 2: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif ch.name == "p":
                p = newp(); runs(p, ch, {})
                if "caption" in (ch.get("class") or []): p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif ch.name in ("ul", "ol"):
                num = int(ch.get("start", 1))
                for li in ch.find_all("li", recursive=False):
                    p = newp()
                    p.add_run(f"{num}. " if ch.name == "ol" else "• ")
                    runs(p, li, {}); num += 1
            elif ch.name == "table" and isinstance(container, type(doc)):
                table(container, ch)
            else:
                blocks(container, ch) if isinstance(container, type(doc)) else runs(newp(), ch, {})
        flush()

    def table(d, t):
        grid, placed = {}, []
        for r, tr in enumerate(t.find_all("tr")):
            c = 0
            for cell in tr.find_all(["td", "th"], recursive=False):
                while (r, c) in grid: c += 1
                rs, cs = int(cell.get("rowspan", 1)), int(cell.get("colspan", 1))
                for i in range(rs):
                    for j in range(cs): grid[(r + i, c + j)] = True
                placed.append((r, c, rs, cs, cell)); c += cs
        if not placed: return
        nrows = max(r for r, _ in grid) + 1; ncols = max(c for _, c in grid) + 1
        tb = d.add_table(rows=nrows, cols=ncols); tb.style = "Table Grid"
        for r, c, rs, cs, cell in placed:
            a = tb.cell(r, c)
            if rs > 1 or cs > 1:
                a = a.merge(tb.cell(min(r + rs, nrows) - 1, min(c + cs, ncols) - 1))
            blocks(a, cell, first_par=a.paragraphs[0])
            if cell.name == "th":
                for p in a.paragraphs:
                    for run in p.runs: run.bold = True
        d.add_paragraph()

    blocks(doc, soup)
    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workdir")
    ap.add_argument("--out", default="output")
    ap.add_argument("--force", action="store_true", help="собрать даже при недостающих страницах")
    ap.add_argument("--no-soffice", action="store_true", help="не использовать LibreOffice")
    a = ap.parse_args()
    work = Path(a.workdir); out = Path(a.out); out.mkdir(exist_ok=True)
    name = work.name
    plan, missing = load_pages(work)
    if missing and not a.force:
        print("НЕ ХВАТАЕТ страниц:", ", ".join(f"p{p:03d}" for p in missing))
        sys.exit(2)
    pages = [(f"p{p:03d}", (work / "html" / f"p{p:03d}.html").read_text(encoding="utf-8"))
             for p in plan["selected"] if p not in missing]
    body, flags = merge(pages)
    html_path = work / f"{name}.html"
    html_path.write_text(TEMPLATE.format(title=name, body=body), encoding="utf-8")

    docx_path = out / f"{name}.docx"
    soffice = None if a.no_soffice else (shutil.which("soffice") or shutil.which("libreoffice"))
    how = "python-docx"
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            r = subprocess.run([soffice, "--headless", "--convert-to", "docx:MS Word 2007 XML",
                                "--outdir", tmp, str(html_path)], capture_output=True, text=True, timeout=300)
            produced = Path(tmp) / f"{name}.docx"
            if r.returncode == 0 and produced.exists():
                shutil.move(str(produced), docx_path); how = "LibreOffice"
    if how == "python-docx":
        to_docx_fallback(body, docx_path)

    fl = out / f"{name}_flags.md"
    fl.write_text(f"# {name}\n\n" + ("\n".join(flags) if flags else "Флагов нет.") + "\n", encoding="utf-8")
    print(f"{docx_path}  ({how}); страниц: {len(pages)}; флагов: {len(flags)}"
          + (f"; НЕ ХВАТАЕТ: {len(missing)}" if missing else ""))


if __name__ == "__main__":
    main()
